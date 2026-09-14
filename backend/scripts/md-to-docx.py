#!/usr/bin/env python3
"""把项目内的 Markdown 文档转换为排版良好的 Word 文档。

支持标题、段落、粗体、有序与无序列表、表格、代码块与水平分隔线，
并针对中文设置了字体（正文宋体、标题黑体、代码 Consolas）。

用法：
    python md-to-docx.py <input.md> <output.docx> [--title "文档标题"]
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

BODY_FONT = "宋体"
HEADING_FONT = "黑体"
CODE_FONT = "Consolas"

HEADING_SIZES = {1: 18, 2: 15, 3: 13.5, 4: 12.5, 5: 12, 6: 11.5}


def set_run_font(run, name, size=None, bold=None, color=None):
    """统一设置中西文字体，避免中文回退到默认字体。"""
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_rich_text(paragraph, text, base_size=None, base_bold=False):
    """解析 **粗体** 与 `行内代码`，其余按普通文本写入。"""
    pattern = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")
    for piece in pattern.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            run = paragraph.add_run(piece[2:-2])
            set_run_font(run, BODY_FONT, base_size, True)
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            run = paragraph.add_run(piece[1:-1])
            set_run_font(run, CODE_FONT, (base_size or 10.5) - 1)
        else:
            run = paragraph.add_run(piece)
            set_run_font(run, BODY_FONT, base_size, base_bold or None)
    return paragraph


def split_table_row(line):
    cells = line.strip().strip("|").split("|")
    return [cell.strip() for cell in cells]


def is_separator_row(line):
    return bool(re.fullmatch(r"\|[\s:|-]+\|", line.strip()))


def add_table(document, rows):
    header, body = rows[0], rows[1:]
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, text in enumerate(header):
        cell = table.rows[0].cells[index]
        cell.text = ""
        add_rich_text(cell.paragraphs[0], text, base_size=10, base_bold=True)
    for row in body:
        cells = table.add_row().cells
        for index in range(len(header)):
            cells[index].text = ""
            add_rich_text(cells[index].paragraphs[0], row[index] if index < len(row) else "", base_size=10)
    document.add_paragraph()


def convert(markdown_path, output_path, title=None):
    lines = Path(markdown_path).read_text(encoding="utf-8").splitlines()
    document = Document()

    style = document.styles["Normal"]
    style.font.name = BODY_FONT
    style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    style.font.size = Pt(11)

    if title:
        heading = document.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = heading.add_run(title)
        set_run_font(run, HEADING_FONT, 20, True)

    index = 0
    in_code = False
    code_buffer = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        # 代码块
        if stripped.startswith("```"):
            if in_code:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Pt(18)
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(10)
                run = paragraph.add_run("\n".join(code_buffer))
                set_run_font(run, CODE_FONT, 9.5)
                code_buffer = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_buffer.append(line)
            index += 1
            continue

        # 表格
        if stripped.startswith("|") and index + 1 < len(lines) and is_separator_row(lines[index + 1]):
            rows = [split_table_row(stripped)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(split_table_row(lines[index]))
                index += 1
            add_table(document, rows)
            continue

        # 水平线
        if stripped in ("---", "***", "___"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(6)
            run = paragraph.add_run("─" * 40)
            set_run_font(run, BODY_FONT, 10, color=RGBColor(0xAA, 0xAA, 0xAA))
            index += 1
            continue

        # 标题
        heading_match = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2).strip()
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
            paragraph.paragraph_format.space_after = Pt(6)
            run = paragraph.add_run(text)
            set_run_font(run, HEADING_FONT, HEADING_SIZES.get(level, 12), True)
            index += 1
            continue

        # 引用
        if stripped.startswith(">"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(18)
            add_rich_text(paragraph, stripped.lstrip("> ").strip(), base_size=10)
            index += 1
            continue

        # 无序列表
        bullet_match = re.match(r"^[-*]\s+(.*)$", stripped)
        if bullet_match:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(2)
            add_rich_text(paragraph, bullet_match.group(1))
            index += 1
            continue

        # 有序列表
        ordered_match = re.match(r"^\d+\.\s+(.*)$", stripped)
        if ordered_match:
            paragraph = document.add_paragraph(style="List Number")
            paragraph.paragraph_format.space_after = Pt(2)
            add_rich_text(paragraph, ordered_match.group(1))
            index += 1
            continue

        # 空行
        if not stripped:
            index += 1
            continue

        # 普通段落
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.first_line_indent = Pt(22)
        add_rich_text(paragraph, stripped)
        index += 1

    document.save(output_path)
    return output_path


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        return 1
    source = sys.argv[1]
    target = sys.argv[2]
    title = None
    if "--title" in sys.argv:
        title = sys.argv[sys.argv.index("--title") + 1]
    convert(source, target, title)
    print(f"generated: {target}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
